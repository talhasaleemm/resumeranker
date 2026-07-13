import fitz
import docx
from pathlib import Path

SAMPLE_DIR = Path("tests/sample_resumes")
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

def generate_multicolumn():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "EXPERIENCE\nCompany A\nSoftware Engineer\n2020-2023\n\nEDUCATION\nBS Computer Science\n2019", fontsize=12)
    page.insert_text((300, 50), "SKILLS\nPython\nJava\nC++\nDocker\nAWS\n\nPROJECTS\nBuilt a cool app\n", fontsize=12)
    doc.save(SAMPLE_DIR / "resume_multicolumn.pdf")

def generate_table_based():
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).text = "EXPERIENCE"
    table.cell(0,1).text = "Company A, 2020-2023. Software Engineer doing Python."
    table.cell(1,0).text = "SKILLS"
    table.cell(1,1).text = "Python, Docker, Kubernetes"
    doc.save(SAMPLE_DIR / "resume_table.docx")

def generate_scanned():
    # Create a temporary PDF with text
    temp = fitz.open()
    temp_page = temp.new_page()
    temp_page.insert_text((50, 50), "Scanned Resume\njohn@scanned.com\nSKILLS: Python, Java", fontsize=12)
    pix = temp_page.get_pixmap()
    
    # Create the actual scanned PDF from the image
    doc = fitz.open()
    page = doc.new_page(width=pix.width, height=pix.height)
    page.insert_image(page.rect, pixmap=pix)
    doc.save(SAMPLE_DIR / "resume_scanned.pdf")

def generate_missing_sections():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Jane Doe\njane@example.com\n\nI am a hard worker. I like coding in Python.", fontsize=12)
    doc.save(SAMPLE_DIR / "resume_missing_sections.pdf")

if __name__ == "__main__":
    generate_multicolumn()
    generate_table_based()
    generate_scanned()
    generate_missing_sections()
    print("Generated edge case resumes.")
